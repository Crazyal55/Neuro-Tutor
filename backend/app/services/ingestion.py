"""Document parsing, chunking, and ingestion into Qdrant."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List

from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.vector_db import delete_by_material_id, upsert_chunks
from app.models.subjects import Material
from app.services.embeddings import embed_texts

logger = logging.getLogger(__name__)

CHUNK_WORDS = 400
OVERLAP_WORDS = 50

SUPPORTED_EXTENSIONS = {
    ".pptx": "pptx",
    ".pdf": "pdf",
    ".docx": "docx",
    ".md": "md",
    ".txt": "txt",
}


@dataclass
class TextBlock:
    text: str
    locator: str


def detect_file_type(filename: str) -> str | None:
    return SUPPORTED_EXTENSIONS.get(Path(filename).suffix.lower())


def chunk_text_blocks(blocks: List[TextBlock]) -> List[TextBlock]:
    """Split text blocks into overlapping word-based chunks."""
    chunks: List[TextBlock] = []

    for block in blocks:
        words = block.text.split()
        if not words:
            continue

        start = 0
        while start < len(words):
            end = min(start + CHUNK_WORDS, len(words))
            chunk_words = words[start:end]
            chunks.append(
                TextBlock(
                    text=" ".join(chunk_words),
                    locator=block.locator,
                )
            )
            if end >= len(words):
                break
            start = max(end - OVERLAP_WORDS, start + 1)

    return chunks


def parse_file(file_path: Path, file_type: str) -> List[TextBlock]:
    if file_type == "txt" or file_type == "md":
        content = file_path.read_text(encoding="utf-8", errors="ignore")
        heading_blocks = _split_markdown_sections(content)
        return heading_blocks or [TextBlock(text=content, locator="document")]

    if file_type == "pdf":
        return _parse_pdf(file_path)

    if file_type == "docx":
        return _parse_docx(file_path)

    if file_type == "pptx":
        return _parse_pptx(file_path)

    raise ValueError(f"Unsupported file type: {file_type}")


def _split_markdown_sections(content: str) -> List[TextBlock]:
    sections = re.split(r"^(#{1,3}\s+.+)$", content, flags=re.MULTILINE)
    if len(sections) <= 1:
        return []

    blocks: List[TextBlock] = []
    heading = "section"
    for index, part in enumerate(sections):
        part = part.strip()
        if not part:
            continue
        if part.startswith("#"):
            heading = part.lstrip("#").strip()
        else:
            blocks.append(TextBlock(text=part, locator=heading))
    return blocks


def _parse_pdf(file_path: Path) -> List[TextBlock]:
    import fitz

    blocks: List[TextBlock] = []
    with fitz.open(file_path) as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                blocks.append(TextBlock(text=text, locator=f"p. {page_number}"))
    return blocks


def _parse_docx(file_path: Path) -> List[TextBlock]:
    from docx import Document

    document = Document(str(file_path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    if not paragraphs:
        return []
    return [TextBlock(text="\n".join(paragraphs), locator="document")]


def _parse_pptx(file_path: Path) -> List[TextBlock]:
    from pptx import Presentation

    presentation = Presentation(str(file_path))
    blocks: List[TextBlock] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Speaker notes: {notes}")
        if texts:
            blocks.append(
                TextBlock(
                    text="\n".join(texts),
                    locator=f"slide {slide_number}",
                )
            )
    return blocks


def ingest_material(db: Session, material_id: str) -> None:
    """Parse, chunk, embed, and upsert a material. Updates status in DB."""
    material = db.query(Material).filter(Material.id == material_id).first()
    if material is None:
        return

    material.status = "processing"
    material.error = None
    db.commit()

    try:
        delete_by_material_id(material.id)
        file_path = Path(material.file_path)
        blocks = parse_file(file_path, material.file_type)
        chunks = chunk_text_blocks(blocks)

        payloads = [
            {
                "material_id": material.id,
                "subject_id": material.subject_id,
                "material_kind": material.material_kind,
                "source_filename": material.filename,
                "locator": chunk.locator,
                "text": chunk.text,
            }
            for chunk in chunks
        ]

        vectors = embed_texts([payload["text"] for payload in payloads])
        if payloads and not upsert_chunks(payloads, vectors):
            raise RuntimeError("Vector store unavailable")

        material.status = "ready"
        material.chunk_count = len(payloads)
        material.error = None
        db.commit()
    except Exception as error:
        logger.exception("Failed to ingest material %s", material_id)
        material.status = "failed"
        material.error = str(error)
        material.chunk_count = 0
        db.commit()
